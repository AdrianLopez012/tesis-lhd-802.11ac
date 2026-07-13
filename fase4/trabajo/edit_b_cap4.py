# -*- coding: utf-8 -*-
# ETAPA B — Redacción del Capítulo 4 completo (estructura Lección 09 TEL147)
# Inserta secciones 4.1-4.11 antes de CONCLUSIONES. Números desde los CSVs.
import docx, sys, csv, glob, os
import numpy as np
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.stdout.reconfigure(encoding="utf-8")

DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
RES = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\cap3\simulacion_ns3\results"
GS  = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\cap3\simulacion_ns3\graficas_simulacion"

# ---- KPIs desde los CSVs ----
def flows(path):
    return {r["name"]: {k: (float(v) if k != "name" else v) for k, v in r.items()}
            for r in csv.DictReader(open(path))}
def agg(name, key):
    v = [flows(p)[name][key] for p in sorted(glob.glob(os.path.join(RES, "principal_s*_v3_flow_stats.csv")))]
    a = np.array(v); return a.mean()
def metr(pat, key):
    v = []
    for p in sorted(glob.glob(os.path.join(RES, pat))):
        for r in csv.DictReader(open(p)):
            if r.get("metric") == key: v.append(float(r.get("value_ms") or r.get("value")))
    return v

owd = agg("Comandos","owd_ms"); plr = agg("Comandos","plr_pct")
e2e = agg("Video","e2e_ms"); jit = agg("Video","jitter_p95_ms")
gp  = agg("Video","goodput_mbps"); plv = agg("Video","plr_pct")
rtt = np.mean(metr("principal_s*_v3_rtt.csv","rtt_media"))
disp = np.mean(metr("principal_s*_v3_disponibilidad.csv","disponibilidad_pct"))
ho_max = max(metr("handover*_v3_handover.csv","max"))
ev = flows(os.path.join(RES,"estres_video_v3_flow_stats.csv"))

d = docx.Document(DOC)
assert d.paragraphs[820].text.strip().startswith("Capítulo 4"), d.paragraphs[820].text[:50]
assert d.paragraphs[822].text.strip() == "CONCLUSIONES", d.paragraphs[822].text[:50]
ref = d.paragraphs[822]          # todo se inserta ANTES de CONCLUSIONES
ph  = d.paragraphs[821]          # placeholder a reemplazar

CAPTION_STYLE = d.paragraphs[798].style   # estilo de los captions existentes
TABLE_STYLE   = d.tables[22].style        # estilo de las tablas existentes

# ---- helpers de inserción ----
def para(txt, style=None):
    p = ref.insert_paragraph_before(txt)
    if style: p.style = style
    return p
def h2(txt): return para(txt, "Heading 2")
def h3(txt): return para(txt, "Heading 3")
def caption(txt):
    p = ref.insert_paragraph_before(txt); p.style = CAPTION_STYLE; return p
def tabla(rows):
    anchor = ref.insert_paragraph_before("")
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = TABLE_STYLE
    for i, fila in enumerate(rows):
        for j, val in enumerate(fila):
            t.cell(i, j).text = val
    anchor._p.addnext(t._element)
    return t
def figura(path, ancho_cm=15.5):
    p = ref.insert_paragraph_before("")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(ancho_cm))
    return p

# ---- prefacio (reemplaza el placeholder) ----
ph.runs[0].text = (
"Este capítulo valida la solución diseñada en el Capítulo 3 frente a criterios técnicos, económicos, ambientales, "
"de seguridad de la información, regulatorios, sociales y éticos. El propósito no es describir nuevamente la solución, "
"sino demostrar con evidencia que resulta idónea para el desafío de ingeniería formulado, dentro del alcance declarado "
"de un trabajo de gabinete. La evidencia técnica proviene de la batería de simulación del prototipo v3-REAL; los demás "
"criterios se evalúan a partir de la documentación del caso de estudio, la normativa aplicable y las consideraciones "
"no técnicas desarrolladas en la sección 1.5.")
for r in ph.runs[1:]: r.text = ""

# ================= 4.1 =================
h2("4.1 Metodología de validación")
para("La validación de idoneidad adopta un enfoque multicriterio. Cada dimensión se evalúa con el instrumento que corresponde a su naturaleza: "
"la dimensión técnica, mediante el contraste cuantitativo entre los indicadores medidos en la batería de simulación y los requisitos no funcionales "
"definidos en el Capítulo 3; la económica, mediante el análisis de la estructura de costos e ingresos evitados, sin cuantificación monetaria que "
"exigiría cotizaciones vigentes de fabricantes; la ambiental, mediante una revisión cualitativa de ciclo de vida inspirada en el enfoque de la "
"norma ISO 14040; y las dimensiones regulatoria, social y ética, mediante el contraste con la normativa peruana aplicable y con los compromisos "
"declarados en la sección 1.5.")
para("La evidencia técnica primaria es la batería de dieciocho corridas del prototipo ns-3 v3-REAL descrita en la sección 3.4.7: diez semillas "
"independientes del escenario de operación (reportadas como media e intervalo de confianza al 95 %), una referencia estática, dos escenarios de "
"estrés y cinco semillas del escenario dedicado de traspaso. Se complementa con el presupuesto de enlace, el mapa de cobertura calculado con el "
"mismo modelo de propagación de la simulación y el contraste de coherencia con el reporte TamoGraph. Toda la evidencia es reproducible: los "
"parámetros provienen de una fuente única de configuración y los resultados se conservan en archivos de datos trazables (Anexos C a F).")
para("El alcance de esta validación es el de una tesis basada en prototipado conceptual: demuestra idoneidad de diseño, no certifica una "
"instalación. Las validaciones que requieren campo —medición radioeléctrica punto a punto, integración con equipos reales y pruebas con el "
"vehículo— se identifican expresamente como trabajo posterior.")

# ================= 4.2 =================
h2("4.2 Idoneidad técnica")
h3("4.2.1 Cumplimiento de requisitos funcionales y no funcionales")
para(f"La solución cumple la totalidad de los indicadores definidos en las Tablas 11 y 12, con márgenes amplios. El retardo medio de comandos "
f"({owd:.2f} ms) utiliza menos del 16 % del presupuesto de 20 ms; la latencia extremo a extremo del video ({e2e:.2f} ms, incluidos los 35 ms de "
f"códec) utiliza menos de la cuarta parte del umbral de 150 ms; y el jitter P95 ({jit:.2f} ms) se mantiene dos órdenes de magnitud por debajo del "
f"límite de 10 ms. El goodput útil de video ({gp:.2f} Mbps de payload) supera el mínimo de 38 Mbps, y las pérdidas de comandos ({plr:.2f} %) "
f"satisfacen en promedio incluso la meta estricta de 0.1 % definida para enlace estable.")
para(f"Los indicadores de continuidad completan el cuadro: el lazo de control presenta un RTT de {rtt:.2f} ms frente al requisito de 40 ms; el "
f"peor traspaso medido en el escenario dedicado fue de {ho_max:.2f} ms frente al requisito de 150 ms; y la disponibilidad radioeléctrica del "
f"recorrido fue del {disp:.0f} % frente al 99.9 % requerido. En conjunto, la red diseñada hace lo que debe hacer: sostiene la conducción remota "
"con video fluido, comandos oportunos y enlace continuo a lo largo del ciclo real de operación del LHD.")
para("La Figura 11 sitúa estos resultados frente a los requisitos de la tesis y frente a valores de referencia de la literatura de teleoperación, "
"mostrando que los márgenes obtenidos no son ajustados sino holgados en todos los indicadores.")
figura(os.path.join(GS, "comparacion_kpis.png"))
caption("Figura 11. Indicadores medidos frente a los requisitos de la tesis y referencias de aplicación.")
caption("Fuente: Elaboración propia a partir de los resultados de la batería de simulación v3-REAL.")

h3("4.2.2 Compatibilidad con estándares técnicos")
para("La solución se construye íntegramente sobre estándares abiertos y de adopción industrial masiva: IEEE 802.11ac (Wave 2) en 5 GHz con "
"canal de 40 MHz y MIMO 2×2 para el acceso inalámbrico; priorización de tráfico conforme a IEEE 802.11e mediante las categorías de acceso WMM; "
"y transporte Ethernet/IEEE 802.3 sobre el anillo de fibra óptica documentado del nivel. Los equipos considerados (nodos Hawk y Cardinal con "
"mesh de capa 2) operan estas tecnologías de fábrica, por lo que el diseño no depende de desarrollos a medida ni de protocolos experimentales.")

h3("4.2.3 Solidez y límites de la validación por simulación")
para(f"La validación técnica es sólida en tres sentidos. Primero, es estadística: el escenario de operación se ejecutó con diez semillas "
f"independientes y los resultados se reportan con su intervalo de confianza. Segundo, es adversa: los escenarios de estrés exploran condiciones "
f"peores que las nominales; al elevar el video a 50 Mbps el sistema mantiene los criterios pero el retardo de comandos se multiplica "
f"(OWD de {ev['Comandos']['owd_ms']:.1f} ms), lo que demuestra que el dimensionamiento a 40 Mbps conserva una reserva verificable y que la "
"priorización WMM protege al tráfico crítico degradando primero la latencia y no las pérdidas. Tercero, es reproducible: geometría, recorrido y "
"parámetros radioeléctricos se generan desde una fuente única de datos y la batería completa puede repetirse con un único script.")
para("Los límites también son claros y se declaran en la Tabla 23: el modelo de propagación es determinista (la variabilidad por desvanecimiento "
"se analiza por separado en el mapa de cobertura y el contraste TamoGraph), el roaming del mesh industrial se emula mediante histéresis de "
"asociación, y no se modelan interferencia externa, fallas de infraestructura ni múltiples vehículos. Ninguna de estas simplificaciones invalida "
"la conclusión de diseño; delimitan lo que la fase de campo deberá confirmar.")

# ================= 4.3 =================
h2("4.3 Idoneidad económica y financiera")
h3("4.3.1 Estructura de costos")
para("La inversión requerida es incremental respecto de la infraestructura ya documentada del nivel: el anillo de fibra óptica, los nodos de red "
"y parte del equipamiento de comunicaciones existen en el caso de estudio, de modo que el costo de capital se concentra en los equipos de acceso "
"inalámbrico de la zona de teleoperación y en su integración. La Tabla 25 organiza la estructura de costos por categoría; su cuantificación "
"monetaria requiere cotizaciones vigentes de los fabricantes y se identifica como paso previo a la decisión de inversión, en coherencia con el "
"enfoque de evaluación descrito en la sección 1.5.4.")
tabla([
 ["Categoría","Componentes principales","Naturaleza"],
 ["Equipamiento de acceso","5 nodos Hawk, 7 nodos Cardinal, radio embarcado con antena HELI-40","CAPEX"],
 ["Integración y red","Puertos y switches industriales, gateway del mesh, energización PoE, cableado local","CAPEX"],
 ["Instalación","Montaje en galería, obra menor, comisionamiento y pruebas de aceptación","CAPEX"],
 ["Estación de teleoperación","Consola de operador, pantallas, controles y software de operación","CAPEX"],
 ["Operación y mantenimiento","Inspección de antenas y conectores, firmware, monitoreo de KPIs, repuestos","OPEX"],
 ["Capacitación","Formación de operadores y personal de mantenimiento OT/TI","OPEX"],
])
caption("Tabla 25. Estructura de costos de implementación y operación de la solución.")

h3("4.3.2 Análisis de retorno y eficiencia")
para("El retorno de la inversión opera por tres mecanismos. El primero es la reducción del riesgo de accidentes en el frente de producción: "
"cada evento evitado representa, además del valor humano primordial, la eliminación de paradas de operación, investigaciones y sobrecostos "
"asociados. El segundo es la continuidad operativa: la teleoperación permite extraer mineral en ventanas en las que hoy el personal no puede "
"ingresar —por ejemplo, tras voladuras o en condiciones de terreno observadas—, convirtiendo tiempo muerto en tiempo productivo. El tercero es "
"la eficiencia de infraestructura: al reutilizar el backbone óptico existente, la inversión evita duplicar red de transporte. La cuantificación "
"del retorno requiere datos operativos de la mina (horas de parada evitables, tonelaje por hora, costos de operación) y se recomienda como "
"análisis conjunto con el área de planeamiento.")

h3("4.3.3 Viabilidad financiera contextual")
para("En el contexto de una operación polimetálica subterránea de la escala de Cerro Lindo, el orden de magnitud de la inversión —una docena de "
"equipos de acceso industrial más su integración— es compatible con los presupuestos anuales de tecnología operacional del sector. La solución "
"no exige licenciamiento de espectro ni infraestructura civil mayor, y su despliegue puede ser gradual: un corredor piloto de teleoperación "
"primero, con expansión por frentes según resultados. Esta gradualidad reduce el riesgo financiero de la adopción.")

# ================= 4.4 =================
h2("4.4 Idoneidad ambiental y de sostenibilidad")
h3("4.4.1 Análisis de ciclo de vida")
para("En la fase de diseño, el trabajo por simulación evita despliegues físicos de prueba y su huella asociada. En la implementación, el "
"equipamiento es de grado industrial con vida útil prolongada y consumo eléctrico alimentado por PoE, marginal frente al consumo del parque de "
"maquinaria de una operación subterránea. En la operación, la teleoperación abre la posibilidad de optimizar los ciclos del LHD y reducir el "
"consumo específico por tonelada extraída, beneficio potencial que deberá medirse en la práctica. En el retiro, los equipos electrónicos deben "
"gestionarse conforme a la normativa de residuos de aparatos eléctricos y electrónicos, mediante los programas de manejo del titular minero.")
para("El alineamiento con los Objetivos de Desarrollo Sostenible descrito en la sección 3.5.3 (ODS 3, 8 y 9) se complementa aquí con el ODS 12 "
"(producción responsable) en lo relativo al ciclo de vida del equipamiento. El impacto ambiental neto de la solución es reducido y su "
"contribución principal es social: retirar personas de zonas de riesgo.")

# ================= 4.5 =================
h2("4.5 Seguridad y privacidad de la información")
para("La red transporta dos activos de información sensibles: los comandos de control del vehículo y el video del frente de operación. Para los "
"primeros, el diseño contempla autenticación robusta de equipos, cifrado del enlace inalámbrico conforme a las capacidades de los equipos "
"industriales considerados, segmentación del tráfico de control respecto de otras redes de la mina y registro de eventos para auditoría. Un "
"comando falsificado o alterado tendría consecuencias físicas, por lo que la seguridad de la información es aquí un requisito de seguridad "
"operacional y no un complemento.")
para("Para el video, que puede captar a trabajadores en el entorno de la operación, aplica la Ley N.° 29733 de Protección de Datos Personales: "
"el tratamiento debe limitarse a la finalidad operativa declarada, con acceso restringido, retención mínima necesaria y comunicación "
"transparente al personal, en línea con los compromisos de privacidad asumidos en la sección 1.5.2. La simulación no implementa estos "
"mecanismos; los define como requisitos de la implementación.")

# ================= 4.6 =================
h2("4.6 Idoneidad regulatoria y legal")
para("El uso de la banda de 5 GHz se enmarca en las condiciones de uso libre establecidas por la normativa peruana de telecomunicaciones para "
"esta banda, por lo que la solución no requiere licenciamiento de espectro; adicionalmente, al operar en interior mina, el entorno confinado "
"minimiza tanto la interferencia recibida como la emitida hacia terceros. En materia de seguridad minera, el diseño se alinea con el Reglamento "
"de Seguridad y Salud Ocupacional en Minería (Decreto Supremo N.° 024-2016-EM y sus modificatorias), cuyo espíritu de jerarquía de controles "
"privilegia eliminar la exposición al peligro: la teleoperación es precisamente un control de ingeniería que retira al trabajador del frente.")
para("La instalación deberá cumplir además los estándares internos del titular minero para trabajos en interior mina y las consideraciones "
"sobre atmósferas y condiciones del entorno tratadas en la sección 1.5.1. La Tabla 26 resume la matriz normativa aplicable y su implicancia "
"para el diseño.")
tabla([
 ["Ámbito","Instrumento","Implicancia para la solución"],
 ["Espectro radioeléctrico","Normativa MTC de bandas de uso libre (5 GHz)","Operación sin licenciamiento; potencias dentro de los límites permitidos."],
 ["Seguridad y salud minera","D.S. 024-2016-EM y modificatorias","La teleoperación actúa como control de ingeniería; requiere procedimientos y validación antes de operar con personal."],
 ["Protección de datos","Ley N.° 29733 y su reglamento","Tratamiento del video con finalidad declarada, acceso restringido y retención mínima."],
 ["Estándares técnicos","IEEE 802.11ac / 802.11e / 802.3","Interoperabilidad y soporte industrial de largo plazo."],
 ["Normas internas","Estándares del titular minero","Permisos de trabajo, homologación de equipos e integración con SSO y OT/TI."],
])
caption("Tabla 26. Matriz normativa aplicable y su implicancia para el diseño.")

# ================= 4.7 =================
h2("4.7 Idoneidad frente a la salud, el bienestar y el bien común")
para("El beneficio central de la solución es de salud ocupacional: retira al operador de un entorno con riesgo de desprendimientos, polvo, "
"gases, ruido, vibración y tránsito de maquinaria pesada, y lo reubica en una estación de control con condiciones ergonómicas controladas. "
"Este beneficio es directo, medible en exposición evitada, y constituye la justificación primaria del proyecto declarada desde el Capítulo 1.")
para("El diseño incorpora una salvaguarda explícita contra el riesgo de falsa seguridad: los requisitos definen un modo seguro ante pérdida de "
"comunicación, pérdida de video o degradación persistente del enlace, y la validación técnica se declara como evidencia de diseño, no como "
"certificación de seguridad funcional. El bien común se expresa también a escala sectorial: demostrar que una red basada en estándares "
"abiertos puede sostener teleoperación en block caving aporta un camino replicable para la seguridad laboral de la minería subterránea peruana.")

# ================= 4.8 =================
h2("4.8 Idoneidad frente a la cultura y la sociedad")
para("La teleoperación transforma el perfil ocupacional del operador de LHD: de la cabina en el frente a una consola de control. Conforme a los "
"compromisos de la sección 1.5.2, esta transición debe gestionarse con reubicación responsable y capacitación, convirtiendo el cambio "
"tecnológico en desarrollo de competencias técnicas para el personal actual en lugar de sustitución. La solución genera además demanda de "
"nuevos roles locales de soporte OT/TI y mantenimiento de redes industriales.")
para("La aceptación social del cambio requiere transparencia: comunicar al personal la finalidad de seguridad del sistema, los alcances del "
"registro de video y los criterios de operación. La tecnología es culturalmente viable precisamente porque su propósito declarado —proteger la "
"vida del trabajador minero— conecta con la prioridad más compartida del sector.")

# ================= 4.9 =================
h2("4.9 Idoneidad ética, moral y deontológica")
para("Las decisiones éticas del trabajo son verificables en el propio documento. Primero, la veracidad: los resultados se reportan con sus "
"supuestos, intervalos de confianza y limitaciones declaradas, y en ningún punto la simulación se presenta como certificación de una red "
"desplegada; este deber de honestidad profesional es exigido por el código deontológico del Colegio de Ingenieros del Perú. Segundo, la "
"prudencia: los criterios de aceptación distinguen la meta estricta del criterio operacional, y las funciones de seguridad se condicionan a "
"validación en campo. Tercero, la responsabilidad sobre las personas: la finalidad del sistema es reducir exposición, su despliegue exige un "
"modo seguro ante fallas, y el tratamiento del video respeta la privacidad del personal.")
para("El dilema ético central —automatizar tareas que hoy realizan personas— se resuelve en este caso a favor de la intervención: el puesto "
"que se transforma es precisamente el que concentra la exposición al riesgo, y la gestión del cambio comprometida en la sección 1.5 preserva "
"el empleo mediante reconversión. La solución es, por diseño y por gestión, éticamente defendible.")

# ================= 4.10 =================
h2("4.10 Síntesis de hallazgos y recomendaciones")
para("La validación multicriterio arroja un resultado consistente: la solución es técnicamente idónea con márgenes holgados y evidencia "
"estadística; económicamente razonable por su carácter incremental sobre infraestructura existente, a falta de la cuantificación con "
"cotizaciones vigentes; ambientalmente de impacto reducido; conforme al marco regulatorio peruano sin requerir licenciamiento de espectro; y "
"socialmente y éticamente sólida, con la seguridad del trabajador como propósito y con salvaguardas explícitas contra la falsa seguridad.")
para("De la validación se derivan cinco recomendaciones para la fase siguiente: (1) realizar la campaña de medición radioeléctrica punto a "
"punto en el nivel para calibrar el modelo con datos propios; (2) obtener cotizaciones de fabricantes y ejecutar el análisis costo-beneficio "
"con datos operativos de la mina; (3) implementar un piloto controlado en el corredor de teleoperación con equipos reales antes de operación "
"regular; (4) definir la política de tratamiento del video conforme a la Ley N.° 29733 antes de instalar cámaras; y (5) desarrollar el plan de "
"capacitación y reconversión del personal operador en paralelo al despliegue técnico.")

# ================= 4.11 =================
h2("4.11 Conclusiones del capítulo")
para("La solución diseñada supera la validación multicriterio dentro del alcance declarado. La evidencia técnica es la más fuerte: todos los "
"indicadores de desempeño cumplen sus umbrales con margen y respaldo estadístico sobre la geometría real de la zona de producción. Las "
"dimensiones económica, ambiental, regulatoria, social y ética no presentan impedimentos y sí beneficios netos, con condiciones de "
"implementación identificadas y acotadas. Se concluye que la red IEEE 802.11ac diseñada es idónea para la teleoperación del LHD en el nivel "
"NV1640 y que el camino hacia el despliegue está definido por las validaciones de campo recomendadas, no por incertidumbres de diseño.")

d.save(DOC)
print("Capítulo 4 insertado: 11 secciones, 2 tablas (25, 26), 1 figura (11)")
