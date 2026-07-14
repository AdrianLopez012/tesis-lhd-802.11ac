# -*- coding: utf-8 -*-
# Convierte los 5 captions manuales (Figuras 9-13) a campos SEQ de Word, para
# que TODO el documento use numeración automática consistente. Luego revisa
# las referencias de texto del Anexo A.
import docx, re, sys
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)

def a_seq(p, numero, resto):
    """convierte el caption a: 'Figura ' + [SEQ Figura] + resto"""
    # limpiar runs
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r1 = p.add_run("Figura ")
    fld = p._element.makeelement(qn("w:fldSimple"), {qn("w:instr"): r" SEQ Figura \* ARABIC "})
    rnum = p._element.makeelement(qn("w:r"), {})
    tnum = p._element.makeelement(qn("w:t"), {})
    tnum.text = str(numero)
    rnum.append(tnum)
    fld.append(rnum)
    r1._element.addnext(fld)
    p.add_run(resto)

pat = re.compile(r"^Figura (\d+)\.(.*)$", re.S)
n = 0
for p in d.paragraphs:
    if p.style.name != "Caption":
        continue
    m = pat.match(p.text.strip())
    if not m:
        continue
    num, resto = int(m.group(1)), "." + m.group(2)
    a_seq(p, num, resto)
    n += 1
    print(f"caption -> SEQ: Figura {num}")

# referencias de texto en el Anexo A que hayan quedado corridas por el +2:
# el Anexo A (tras 'Anexo A.') referencia sus figuras; con SEQ serán 14-18.
inA = False
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith("Anexo A."): inA = True
    elif t.startswith("Anexo B."): inA = False
    if inA and re.search(r"Figura \d+", t) and not t.startswith("Figura") and p.runs:
        print(f"REF texto AnexoA: {t[:90]}")

d.save(DOC)
print(f"{n} captions convertidos a SEQ")
