# -*- coding: utf-8 -*-
# ETAPA D3 — limpiar residuos v8.2 en captions, tablas y anexo H; revisar RESUMEN
import docx, sys
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)

def rep_para(p, viejo=None, nuevo_full=None, subs=None):
    txt = p.text
    if subs:
        for a, b in subs: txt = txt.replace(a, b)
    else:
        txt = nuevo_full
    p.runs[0].text = txt
    for r in p.runs[1:]: r.text = ""

# RESUMEN y ABSTRACT: mostrar la mención v8 para decidir
for i in (12, 15):
    t = d.paragraphs[i].text
    pos = t.find("v8")
    print(f"[{i}] ...{t[max(0,pos-80):pos+60]}..." if pos >= 0 else f"[{i}] sin v8")

SUBS = [
 ("prototipo ns-3 v8.2", "prototipo ns-3 v3-REAL"),
 ("simulación ns-3 v8.2", "simulación ns-3 v3-REAL"),
 ("del submodelo v8.2", "del modelo v3-REAL"),
 ("reproducible v8.2", "reproducible v3-REAL"),
 ("resultados v8.2", "resultados de la batería v3-REAL"),
 ("KPIs v8.2", "KPIs de la batería v3-REAL"),
 ("logs v8.2", "logs de la batería v3-REAL"),
 ("Modelo ns-3 v8.2", "Modelo ns-3 v3-REAL"),
 ("v8.2", "v3-REAL"),
]
n = 0
for i, p in enumerate(d.paragraphs):
    if "v8.2" in p.text and p.style.name != "table of figures" and "0001 |" not in p.text[:8]:
        if p.style.name == "CodigoFuente":  # volcado: referencia histórica legítima
            continue
        antes = p.text[:60]
        rep_para(p, subs=SUBS)
        n += 1
        print(f"parrafo [{i}] corregido: {antes!r}")
for ti, t in enumerate(d.tables):
    for row in t.rows:
        for c in row.cells:
            if "v8.2" in c.text:
                for pp in c.paragraphs:
                    if "v8.2" in pp.text and pp.runs:
                        rep_para(pp, subs=SUBS); n += 1
                        print(f"tabla {ti} celda corregida")
d.save(DOC)
print(f"total corregidos: {n}")
