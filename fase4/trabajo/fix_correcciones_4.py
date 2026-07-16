# -*- coding: utf-8 -*-
# Correcciones parte 4:
#  - RECOMENDACIONES: sacar el sectPr incrustado del heading (lo deja huérfano
#    al pie de la página anterior) y ponerlo en un párrafo vacío ANTES.
#  - Todas las secciones ANTES de INTRODUCCIÓN -> lowerRoman sin reinicio
#    (mata cualquier reinicio arábigo en los índices).
#  - titlePg + primera página sin folio para la sección de RECOMENDACIONES.
import docx, sys, copy
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")

F = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(F)
body = d.element.body

def ptext(el):
    if el is None or el.tag != qn('w:p'): return ''
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))

# ---------- RECOMENDACIONES: mover sectPr incrustado ----------
for el in list(body):
    t = ptext(el).strip()
    if t == 'RECOMENDACIONES Y OBSERVACIONES':
        ppr = el.find(qn('w:pPr'))
        spr = ppr.find(qn('w:sectPr')) if ppr is not None else None
        if spr is not None:
            ppr.remove(spr)
            p_break = body.makeelement(qn('w:p'), {})
            nppr = p_break.makeelement(qn('w:pPr'), {})
            nppr.append(spr)
            p_break.append(nppr)
            el.addprevious(p_break)
            print("RECOMENDACIONES: sectPr sacado del heading y puesto antes")
        else:
            print("RECOMENDACIONES: heading ya no tiene sectPr incrustado")
        # quitar pbb sobrante
        if ppr is not None:
            pbb = ppr.find(qn('w:pageBreakBefore'))
            if pbb is not None: ppr.remove(pbb)
        break

d.save(F)
d = docx.Document(F)

# ---------- mapear secciones ----------
def mapa_secciones(doc):
    res, cur = [], []
    for el in doc.element.body:
        if el.tag == qn('w:p'):
            t = ptext(el).strip()
            if t and not cur: cur.append(t[:50])
            ppr = el.find(qn('w:pPr'))
            if ppr is not None and ppr.find(qn('w:sectPr')) is not None:
                res.append(cur[0] if cur else '(vacia)'); cur = []
        elif el.tag == qn('w:tbl'):
            if not cur: cur.append('(tabla)')
    res.append(cur[0] if cur else '(vacia)')
    return res

mapa = mapa_secciones(d)
idx_intro = next(i for i, t in enumerate(mapa) if t.startswith('INTRODUCCIÓN'))
print(f"INTRODUCCIÓN es la sección {idx_intro}")

# ---------- romanos antes de Introducción ----------
for i, sec in enumerate(d.sections):
    if 1 <= i < idx_intro:
        spr = sec._sectPr
        pgn = spr.find(qn('w:pgNumType'))
        if pgn is None:
            pgn = spr.makeelement(qn('w:pgNumType'), {}); spr.append(pgn)
        pgn.set(qn('w:fmt'), 'lowerRoman')
        if pgn.get(qn('w:start')): del pgn.attrib[qn('w:start')]
print(f"secciones 1..{idx_intro-1} en lowerRoman sin reinicio")

# ---------- titlePg para RECOMENDACIONES ----------
for i, sec in enumerate(d.sections):
    txt = (mapa[i] if i < len(mapa) else '') or ''
    if txt.startswith('RECOMENDACIONES'):
        sec.different_first_page_header_footer = True
        fph = sec.first_page_header; fph.is_linked_to_previous = False
        for p in list(fph.paragraphs[1:]): p._p.getparent().remove(p._p)
        if fph.paragraphs:
            for r in list(fph.paragraphs[0].runs): r._r.getparent().remove(r._r)
        fpf = sec.first_page_footer; fpf.is_linked_to_previous = False
        if fpf.paragraphs:
            for r in list(fpf.paragraphs[0].runs): r._r.getparent().remove(r._r)
        print(f"sec{i} (RECOMENDACIONES): primera página sin folio")

# ---------- mapa final ----------
print()
print("=== MAPA FINAL DE SECCIONES ===")
mapa = mapa_secciones(d)
for i, sec in enumerate(d.sections):
    spr = sec._sectPr
    pgn = spr.find(qn('w:pgNumType'))
    fmt = pgn.get(qn('w:fmt')) if pgn is not None else None
    st  = pgn.get(qn('w:start')) if pgn is not None else None
    tp  = spr.find(qn('w:titlePg')) is not None
    txt = mapa[i] if i < len(mapa) else '?'
    print(f"  sec{i}: fmt={fmt} start={st} titlePg={tp} | {txt[:45]!r}")

d.save(F)
print("GUARDADO OK")
