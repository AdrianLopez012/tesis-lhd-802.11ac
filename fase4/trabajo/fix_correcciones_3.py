# -*- coding: utf-8 -*-
# Correcciones parte 3: SECCIONES y NUMERACIÓN.
#  (a) Crear sección propia para Capítulo 4, CONCLUSIONES y RECOMENDACIONES
#      (hoy viven dentro de la sección del Cap. 3 -> por eso muestran folio
#       y Cap4 arranca arriba).
#  (b) Numeración: índices en romanos (quitar reinicios arábigos);
#      INTRODUCCIÓN = arábigo 1.
#  (c) titlePg + encabezado de primera página VACÍO en secciones de capítulo
#      (primera página cuenta pero no muestra folio).
import docx, sys, copy, re
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")

F = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(F)
body = d.element.body

def ptext(el):
    if el is None or el.tag != qn('w:p'): return ''
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))

def tiene_sectpr(el):
    if el is None or el.tag != qn('w:p'): return False
    ppr = el.find(qn('w:pPr'))
    return ppr is not None and ppr.find(qn('w:sectPr')) is not None

# ---------- (a) saltos de sección para Cap4 / CONCLUSIONES / RECOMENDACIONES ----------
NUEVAS_SECCIONES = ['Capítulo 4.', 'CONCLUSIONES', 'RECOMENDACIONES Y OBSERVACIONES']

def sectpr_envolvente(el):
    """el sectPr que aplica al contenido donde vive el (el siguiente en el doc)"""
    sig = el.getnext()
    while sig is not None:
        if sig.tag == qn('w:p'):
            ppr = sig.find(qn('w:pPr'))
            if ppr is not None and ppr.find(qn('w:sectPr')) is not None:
                return ppr.find(qn('w:sectPr'))
        sig = sig.getnext()
    return body.find(qn('w:sectPr'))     # sectPr final del body

insertados = 0
for el in list(body):
    t = ptext(el).strip()
    if not any(t.startswith(x) for x in NUEVAS_SECCIONES): continue
    prev = el.getprevious()
    if tiene_sectpr(prev): continue                     # ya hay salto justo antes
    if el.find(qn('w:pPr')) is not None and el.find(qn('w:pPr')).find(qn('w:sectPr')) is not None:
        continue                                        # el propio heading es un fin de sección
    # crear párrafo vacío con sectPr copiado de la sección envolvente
    spr_orig = sectpr_envolvente(el)
    spr = copy.deepcopy(spr_orig)
    # el salto es de página nueva y sin reinicio de numeración
    tp = spr.find(qn('w:type'))
    if tp is None:
        tp = spr.makeelement(qn('w:type'), {}); spr.insert(0, tp)
    tp.set(qn('w:val'), 'nextPage')
    pgn = spr.find(qn('w:pgNumType'))
    if pgn is not None and pgn.get(qn('w:start')):
        del pgn.attrib[qn('w:start')]
    p_break = body.makeelement(qn('w:p'), {})
    ppr = p_break.makeelement(qn('w:pPr'), {})
    ppr.append(spr)
    p_break.append(ppr)
    el.addprevious(p_break)
    # quitar el page_break_before del heading (ya no hace falta)
    hppr = el.find(qn('w:pPr'))
    if hppr is not None:
        pbb = hppr.find(qn('w:pageBreakBefore'))
        if pbb is not None: hppr.remove(pbb)
    insertados += 1
    print(f"  [a] sección creada antes de: {t[:50]!r}")
print(f"[a] saltos de sección insertados: {insertados}")

d.save(F)

# ---------- reabrir para que python-docx reconozca las secciones nuevas ----------
d = docx.Document(F)

# mapear sección -> primer texto
def primer_texto_por_seccion(doc):
    res = []
    cur = []
    for el in doc.element.body:
        if el.tag == qn('w:p'):
            t = ptext(el).strip()
            if t and not cur: cur.append(t[:50])
            ppr = el.find(qn('w:pPr'))
            if ppr is not None and ppr.find(qn('w:sectPr')) is not None:
                res.append(cur[0] if cur else '(vacia)')
                cur = []
        elif el.tag == qn('w:tbl'):
            if not cur: cur.append('(tabla)')
    res.append(cur[0] if cur else '(vacia)')
    return res

mapa = primer_texto_por_seccion(d)
print()
print("=== SECCIONES TRAS INSERTAR ===")
for i, s in enumerate(d.sections):
    txt = mapa[i] if i < len(mapa) else '?'
    print(f"  sec{i}: {txt!r}")

# ---------- (b) numeración ----------
ROMANOS = ['ÍNDICE DE TABLAS', 'ÍNDICE DE FIGURAS', 'ÍNDICE DE ECUACIONES']
ARABIGO1 = 'INTRODUCCIÓN'
for i, sec in enumerate(d.sections):
    txt = (mapa[i] if i < len(mapa) else '') or ''
    spr = sec._sectPr
    pgn = spr.find(qn('w:pgNumType'))
    if any(txt.startswith(r) for r in ROMANOS):
        if pgn is None:
            pgn = spr.makeelement(qn('w:pgNumType'), {}); spr.append(pgn)
        pgn.set(qn('w:fmt'), 'lowerRoman')
        if pgn.get(qn('w:start')): del pgn.attrib[qn('w:start')]
        print(f"  [b] sec{i} ({txt[:30]}): lowerRoman, sin reinicio")
    elif txt.startswith(ARABIGO1):
        if pgn is None:
            pgn = spr.makeelement(qn('w:pgNumType'), {}); spr.append(pgn)
        pgn.set(qn('w:fmt'), 'decimal')
        pgn.set(qn('w:start'), '1')
        print(f"  [b] sec{i} (INTRODUCCIÓN): decimal, start=1")

# ---------- (c) titlePg + primera página sin folio en secciones de capítulo ----------
CAPITULOS = ['INTRODUCCIÓN', 'Capítulo 1.', 'Capítulo 2.', 'Capítulo 3.', 'Capítulo 4.',
             'CONCLUSIONES', 'RECOMENDACIONES', 'ANEXOS', 'REFERENCIAS']
n_tp = 0
for i, sec in enumerate(d.sections):
    txt = (mapa[i] if i < len(mapa) else '') or ''
    if any(txt.startswith(c) for c in CAPITULOS):
        sec.different_first_page_header_footer = True
        fph = sec.first_page_header
        fph.is_linked_to_previous = False        # crea encabezado propio VACÍO
        for p in list(fph.paragraphs[1:]):
            p._p.getparent().remove(p._p)
        if fph.paragraphs:
            p0 = fph.paragraphs[0]
            for r in list(p0.runs): r._r.getparent().remove(r._r)
        fpf = sec.first_page_footer
        fpf.is_linked_to_previous = False
        for p in list(fpf.paragraphs[1:]):
            p._p.getparent().remove(p._p)
        if fpf.paragraphs:
            for r in list(fpf.paragraphs[0].runs): r._r.getparent().remove(r._r)
        n_tp += 1
        print(f"  [c] sec{i} ({txt[:35]}): primera página sin folio")
print(f"[c] secciones con primera página sin folio: {n_tp}")

d.save(F)
print("GUARDADO OK")
