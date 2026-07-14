# -*- coding: utf-8 -*-
# Inserta: curva Prx + CDF en Anexo G (tras el patrón), CAPEX/OPEX en 4.3.1,
# y una frase en 3.3.2 que referencia los radios (sin nº de figura).
import docx, sys
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
TRB = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo"
d = docx.Document(DOC)
CAPTION_STYLE = next(p.style for p in d.paragraphs if p.style.name == "Caption")

def fig_seq(ref, img, num_frio, resto_caption, fuente, texto=None, ancho=16.0):
    if texto:
        ref.insert_paragraph_before(texto)
    fp = ref.insert_paragraph_before("")
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run().add_picture(img, width=Cm(ancho))
    cap = ref.insert_paragraph_before("Figura ")
    cap.style = CAPTION_STYLE
    fld = cap._element.makeelement(qn("w:fldSimple"), {qn("w:instr"): r" SEQ Figura \* ARABIC "})
    rn = cap._element.makeelement(qn("w:r"), {}); tn = cap._element.makeelement(qn("w:t"), {})
    tn.text = str(num_frio); rn.append(tn); fld.append(rn)
    cap.runs[0]._element.addnext(fld)
    cap.add_run(resto_caption)
    fu = ref.insert_paragraph_before(fuente)
    fu.style = CAPTION_STYLE

# --- Anexo G: tras el bloque del patrón (ancla = heading Anexo H) ---
iH = next(i for i, p in enumerate(d.paragraphs) if p.text.strip().startswith("Anexo H."))
ref = d.paragraphs[iH]
fig_seq(ref, rf"{TRB}\curva_prx_distancia.png", 21,
        ". Potencia recibida frente a la distancia por ruta de túnel, con las sensibilidades objetivo y el radio de diseño.",
        "Fuente: Elaboración propia con el modelo two-slope de la fuente única de parámetros (MATLAB).",
        texto="La curva siguiente extiende el presupuesto de enlace a distancia continua: el enlace más exigente (Cardinal hacia el LHD) "
        "cruza la sensibilidad de video a 127 m y la de borde a 327 m, más que duplicando el radio de diseño de 60 m.")
fig_seq(ref, rf"{TRB}\cdf_rssi_10seeds.png", 22,
        ". Distribución acumulada del RSSI del enlace asociado y del mejor AP durante el recorrido (diez semillas).",
        "Fuente: Elaboración propia a partir de los registros de posición y asociación de la batería v3-REAL (MATLAB).")

# --- 4.3.1: figura CAPEX/OPEX tras el caption de la Tabla 25 ---
iT25 = next(i for i, p in enumerate(d.paragraphs) if p.text.strip().startswith("Tabla 25. Estructura de costos"))
ref2 = d.paragraphs[iT25 + 1]
fig_seq(ref2, rf"{TRB}\fig_capex_opex.png", 15,
        ". Estructura económica de la solución: inversión, operación y mecanismos de retorno.",
        "Fuente: Elaboración propia.", ancho=15.5)

# --- 3.3.2: frase con los radios (sin número de figura) ---
for p in d.paragraphs:
    if p.text.strip().startswith("La solución IEEE 802.11ac industrial resulta viable"):
        nuevo = p.text.rstrip() + (" El alcance del modelo respalda esta viabilidad: el enlace más exigente mantiene la sensibilidad "
        "de video hasta 127 m y la de borde hasta 327 m de distancia por ruta de túnel —más del doble del radio de diseño de 60 m—, "
        "como se ilustra en la evidencia gráfica del Anexo G.")
        p.runs[0].text = nuevo
        for r in p.runs[1:]: r.text = ""
        break

# --- inventario del Anexo G ---
for p in d.paragraphs:
    if p.text.strip().startswith("— patron_antena_3d.png"):
        p.insert_paragraph_before("— curva_prx_distancia.png: Potencia recibida vs distancia con sensibilidades y radio de diseño (MATLAB).")
        p.insert_paragraph_before("— cdf_rssi_10seeds.png: CDF del RSSI asociado y del mejor AP, diez semillas (MATLAB).")
        p.insert_paragraph_before("— fig_capex_opex.png: Estructura económica de la solución (Figura del Capítulo 4).")
        break
d.save(DOC)
print("3 figuras insertadas + frase 3.3.2 + inventario")
