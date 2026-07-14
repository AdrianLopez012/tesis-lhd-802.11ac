# -*- coding: utf-8 -*-
# Reemplaza la imagen de la Figura 10 (escena 3D v2) y actualiza caption+texto
# preservando el campo SEQ del caption.
import docx, sys, struct
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
IMG = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\escena_mina_3d.png"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

d = docx.Document(DOC)

i_cap = i_txt = None
for i, p in enumerate(d.paragraphs):
    if "Entorno tridimensional" in p.text and p.style.name == "Caption" and i_cap is None:
        i_cap = i
    if p.text.strip().startswith("La Figura 10 presenta el entorno"):
        i_txt = i
assert i_cap and i_txt, (i_cap, i_txt)

# texto previo actualizado
p = d.paragraphs[i_txt]
p.runs[0].text = ("La Figura 10 presenta el entorno en tres dimensiones y consolida visualmente el diseño: el nivel de señal calculado sobre el piso "
"de las galerías con el mismo modelo two-slope de la simulación, el recorrido real del ciclo de operación del LHD, la ubicación etiquetada de los "
"doce AP y el patrón de radiación calculado de una antena helicoidal en modo axial —representativa de la antena embarcada—, ilustrando la relación "
"espacial entre cobertura, infraestructura y movilidad.")
for r in p.runs[1:]: r.text = ""

# caption: preservar 'Figura ' + [SEQ]; editar el run posterior al campo
cap = d.paragraphs[i_cap]
cap.runs[-1].text = (". Entorno tridimensional de la zona de producción: cobertura RSSI calculada, recorrido real del LHD, "
                     "los 12 AP y patrón de radiación de la antena del vehículo.")

# la línea de fuente (párrafo siguiente al caption)
fu = d.paragraphs[i_cap + 1]
if fu.text.strip().startswith("Fuente:"):
    fu.runs[0].text = ("Fuente: Elaboración propia; cobertura con el modelo two-slope de la simulación (fuente única de parámetros) y patrón "
                       "calculado con MATLAB Antenna Toolbox (hélice en modo axial, 5.0 GHz).")
    for r in fu.runs[1:]: r.text = ""

# swap de la imagen: drawing más cercano ANTERIOR al caption
hit = None
for i in range(i_cap - 1, max(i_cap - 6, 0), -1):
    blips = d.paragraphs[i]._element.findall(".//a:blip", NS)
    if blips:
        hit = (i, blips[0].get(R_EMBED)); break
assert hit, "no drawing"
pi, rid = hit
with open(IMG, "rb") as f:
    d.part.related_parts[rid]._blob = f.read()
with open(IMG, "rb") as f:
    head = f.read(24)
w_px, h_px = struct.unpack(">II", head[16:24])
pp = d.paragraphs[pi]
for ext in pp._element.findall(".//wp:extent", NS) + pp._element.findall(".//a:ext", NS):
    cx = int(ext.get("cx"))
    ext.set("cy", str(int(cx * h_px / w_px)))
d.save(DOC)
print(f"figura swapeada (parrafo {pi}, {rid}, {w_px}x{h_px}); caption y textos actualizados")
