# -*- coding: utf-8 -*-
# Inserta el gráfico del RSSI asociado en 3.4.7 (tras el párrafo del roaming)
# con caption SEQ; ajusta la referencia manual del Cap. 4 (Figura 13 -> 14).
import docx, sys
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
IMG = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\grafico_rssi_asociado.png"
d = docx.Document(DOC)

# ancla: el párrafo SIGUIENTE al del roaming en 3.4.7
i_roam = None
for i, p in enumerate(d.paragraphs):
    if p.text.strip().startswith("Los registros de posición y asociación evidencian"):
        i_roam = i; break
assert i_roam, "no hallado parrafo roaming"
ref = d.paragraphs[i_roam + 1]
CAPTION_STYLE = next(p.style for p in d.paragraphs if p.style.name == "Caption")

t = ref.insert_paragraph_before(
"La estabilidad del enlace de servicio se verifica cuantitativamente en la figura siguiente: el RSSI del AP asociado, medido segundo a segundo "
"en las diez semillas, se mantiene durante todo el recorrido por encima de −72.1 dBm en el peor caso —un margen de 9.9 dB sobre el umbral de "
"enlace usable de −82 dBm—, mientras la referencia del mejor AP disponible confirma que la red ofrece siempre una alternativa de mayor señal. "
"El diseño privilegia deliberadamente la estabilidad de asociación sobre la persecución del AP más cercano, evitando los microcortes del "
"ping-pong de reasociaciones; el comportamiento del sistema real de malla, que optimiza rutas dinámicamente, solo puede mejorar este caso "
"conservador.")
fp = ref.insert_paragraph_before("")
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run().add_picture(IMG, width=Cm(16.5))
cap = ref.insert_paragraph_before("Figura ")
cap.style = CAPTION_STYLE
fld = cap._element.makeelement(qn("w:fldSimple"), {qn("w:instr"): r" SEQ Figura \* ARABIC "})
rnum = cap._element.makeelement(qn("w:r"), {}); tnum = cap._element.makeelement(qn("w:t"), {})
tnum.text = "13"; rnum.append(tnum); fld.append(rnum)
cap.runs[0]._element.addnext(fld)
cap.add_run(". RSSI del enlace asociado a lo largo del recorrido (diez semillas) frente a los umbrales de servicio.")
fu = ref.insert_paragraph_before("Fuente: Elaboración propia a partir de los registros de posición y asociación de la batería v3-REAL.")
fu.style = CAPTION_STYLE

# referencia manual del Cap. 4: "La Figura 13 sitúa" -> 14
n = 0
for p in d.paragraphs:
    if "La Figura 13 sitúa" in p.text:
        p.runs[0].text = p.text.replace("La Figura 13 sitúa", "La Figura 14 sitúa")
        for r in p.runs[1:]: r.text = ""
        n += 1
print(f"figura insertada en 3.4.7; refs Cap4 ajustadas: {n}")
d.save(DOC)
