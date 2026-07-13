# -*- coding: utf-8 -*-
# ETAPA D4 — RESUMEN/ABSTRACT: de "resultados preliminares" a consolidados
import docx, sys
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)

def rep(p, subs):
    t = p.text
    for a, b in subs: t = t.replace(a, b)
    p.runs[0].text = t
    for r in p.runs[1:]: r.text = ""

p12 = d.paragraphs[12]
pos = p12.text.find("Los resultados preliminares")
print("RESUMEN antes:", p12.text[pos:pos+220] if pos>=0 else "(no hallado)")
rep(p12, [("Los resultados preliminares de la versión v3-REAL muestran que el escenario principal satisface los criterios",
           "Los resultados de la versión v3-REAL, obtenidos sobre la geometría real de la zona de producción con diez semillas independientes, muestran que el escenario de operación satisface todos los criterios")])

p15 = d.paragraphs[15]
rep(p15, [("Preliminary results from version v8", "x-x-x"),  # no debería quedar
          ("Preliminary results from version v3-REAL show that the principal scenario meets",
           "Results from version v3-REAL, obtained over the real geometry of the production zone with ten independent seeds, show that the operation scenario meets all")])

p1112 = d.paragraphs[1112]
if "v8" in p1112.text:
    print("1112 antes:", p1112.text[:150])
    rep(p1112, [("v8.2", "v3-REAL"), ("v8_2", "v3")])

d.save(DOC)
# verificación final de residuos
d2 = docx.Document(DOC)
resid = [(i, p.text[:70]) for i, p in enumerate(d2.paragraphs)
         if ("v8.2" in p.text or "preliminares de la versión" in p.text.lower())
         and p.style.name not in ("table of figures", "CodigoFuente")]
print("residuales:", resid if resid else "NINGUNO (fuera de índices y código histórico)")
