# -*- coding: utf-8 -*-
# Diagnóstico de tablas: filas vacías (fila de más) y posibles desbordes
import docx, sys
from docx.shared import Emu
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)

# ancho útil de página (sección principal)
s = d.sections[-1]
util = s.page_width - s.left_margin - s.right_margin
print(f"ancho útil de página: {util/914400:.2f} pulgadas\n")

problemas = []
for ti, t in enumerate(d.tables):
    # filas completamente vacías
    for ri, row in enumerate(t.rows):
        if all(not c.text.strip() for c in row.cells):
            problemas.append(f"tabla {ti}: fila {ri} VACÍA ({len(t.rows)} filas totales) — 1a celda fila0: {t.cell(0,0).text.strip()[:40]!r}")
    # ancho declarado de la tabla (suma de anchos de columnas de la fila 0)
    try:
        w = sum((c.width or 0) for c in t.rows[0].cells)
        if w and w > util * 1.02:
            problemas.append(f"tabla {ti}: DESBORDE ancho {w/914400:.2f}in > {util/914400:.2f}in — {t.cell(0,0).text.strip()[:40]!r}")
    except Exception:
        pass
print("\n".join(problemas) if problemas else "sin problemas detectados programáticamente")
print(f"\ntotal tablas: {len(d.tables)}")
