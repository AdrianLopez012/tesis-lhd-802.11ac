# -*- coding: utf-8 -*-
# ETAPA C — CONCLUSIONES y RECOMENDACIONES de la tesis (método Lección 10:
# cada conclusión cierra un objetivo específico con evidencia verificable;
# verbos en presente según observación del asesor).
import docx, sys
sys.stdout.reconfigure(encoding="utf-8")

DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)

# localizar por TEXTO (los índices cambiaron al insertar el Cap. 4)
i_conc = i_conc_ph = i_reco = i_reco_ph = i_refs = None
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t == "CONCLUSIONES": i_conc = i
    elif t.startswith("Las conclusiones finales se elaborarán"): i_conc_ph = i
    elif t == "RECOMENDACIONES Y OBSERVACIONES": i_reco = i
    elif t.startswith("Las recomendaciones finales se consolidarán"): i_reco_ph = i
    elif t == "REFERENCIAS BIBLIOGRÁFICAS" and i_refs is None: i_refs = i
assert None not in (i_conc, i_conc_ph, i_reco, i_reco_ph, i_refs), (i_conc, i_conc_ph, i_reco, i_reco_ph, i_refs)
print(f"CONCLUSIONES en {i_conc}, placeholder {i_conc_ph}; RECOMENDACIONES en {i_reco}, placeholder {i_reco_ph}; REFS en {i_refs}")

def replace(p, txt):
    p.runs[0].text = txt
    for r in p.runs[1:]: r.text = ""

CONCLUSIONES = [
"Las conclusiones siguientes cierran el ciclo argumentativo de la tesis: contrastan cada objetivo específico con la evidencia obtenida y "
"delimitan el alcance de lo demostrado. Todas se sustentan en resultados verificables del trabajo —la batería de simulación del prototipo "
"v3-REAL, el presupuesto de enlace, el mapa de cobertura y el análisis de idoneidad del Capítulo 4— y ninguna extrapola más allá del alcance "
"declarado de un trabajo de gabinete.",

"1. Sobre el objetivo general, el diseño cumple en simulación todas las metas planteadas. En el escenario de operación sobre la geometría real "
"de la zona de producción (diez semillas independientes, media e intervalo de confianza al 95 %), la latencia unidireccional de comandos es de "
"3.04 ± 0.26 ms y la del video de 35.89 ± 0.28 ms incluido el códec —muy por debajo del límite—, el jitter P95 es de 0.28 ms frente a los 10 ms "
"admitidos, la pérdida de comandos (0.06 %) satisface en promedio incluso la meta estricta de 0.1 %, el peor traspaso medido dura 0.91 ms frente "
"a los 150 ms tolerados y la disponibilidad radioeléctrica del recorrido es del 100 %. La red diseñada habilita, dentro del modelo, la "
"teleoperación segura y confiable del LHD.",

"2. Sobre la caracterización del canal (objetivo específico 1), el modelo two-slope calibrado (n1 = 1.9, n2 = 3.4, distancia de quiebre de "
"40 m, pérdidas de sistema de 9.4 dB y penalización NLOS de 10 dB por galería cruzada) describe la propagación en las galerías del caso de "
"estudio de forma coherente con el reporte TamoGraph disponible y con el análisis de ray-tracing por método de imágenes. Los parámetros "
"obtenidos alimentan directamente el presupuesto de enlace y la simulación, con una fuente única de configuración que garantiza consistencia "
"entre todos los artefactos del diseño.",

"3. Sobre la arquitectura (objetivo específico 2), la solución híbrida —anillo de fibra óptica, backbone Ethernet y acceso IEEE 802.11ac con "
"cinco nodos Hawk y siete nodos Cardinal— queda definida sobre las posiciones reales del plano del nivel NV1640, con correspondencia directa "
"entre el modelo y la infraestructura documentada. La comparación de alternativas tecnológicas sustenta la selección de 802.11ac industrial "
"frente a LPWAN, celular privado y óptica inalámbrica para este caso de uso.",

"4. Sobre el dimensionamiento de radio y antenas (objetivo específico 3), el presupuesto de enlace demuestra que el enlace más exigente "
"(Cardinal hacia el LHD con antena HELI-40) conserva un margen de +11.1 dB sobre la sensibilidad de video a la distancia de diseño de 60 m, y "
"que el alcance máximo del modelo (127 m para video y 327 m en borde de celda) más que duplica dicha distancia. La configuración 5 GHz, canal "
"de 40 MHz y MIMO 2×2 sostiene una tasa física observada de hasta 360 Mbps en la simulación, holgada para el tráfico agregado.",

"5. Sobre las políticas de calidad de servicio y movilidad (objetivo específico 4), la priorización IEEE 802.11e/WMM protege eficazmente el "
"tráfico crítico: en el escenario de estrés con video a 50 Mbps, la degradación se manifiesta en la latencia de los comandos (que se duplica "
"pero permanece dentro del umbral) y no en pérdidas. El esquema de movilidad de tipo make-before-break mantiene el enlace durante todo el "
"recorrido; forzados los traspasos duros en el escenario dedicado (cinco semillas), su duración media es de 0.71 ms con peor caso de 0.91 ms.",

"6. Sobre la validación extremo a extremo (objetivo específico 5), la batería de dieciocho corridas —operación multi-semilla, referencia "
"estática, dos escenarios de estrés y traspaso forzado— verifica todos los indicadores clave dentro de las metas, con reproducibilidad "
"completa: geometría, recorrido y parámetros se generan desde una fuente única y la batería se relanza con un único script. La validación es "
"de diseño y no sustituye las pruebas de campo, límite que la tesis declara de forma explícita.",

"7. Sobre la viabilidad técnico-económica (objetivo específico 6), el trabajo entrega la estructura de costos de implementación y operación, "
"la matriz normativa aplicable y los mecanismos de retorno (reducción de exposición, continuidad operativa y reutilización del backbone "
"existente), y demuestra que la inversión es incremental y desplegable de forma gradual. La cuantificación monetaria del CAPEX/OPEX y de los "
"indicadores de decisión requiere cotizaciones vigentes de fabricantes y datos operativos de la mina, y queda establecida como el paso previo "
"a la decisión de inversión.",

"8. Como impacto y aprendizaje, la tesis demuestra que una red basada en estándares abiertos puede sostener la teleoperación de maquinaria "
"pesada en block caving, aportando un camino replicable para retirar personas de las zonas de mayor riesgo de la minería subterránea peruana. "
"Metodológicamente, muestra el valor de una cadena de diseño trazable: una fuente única de parámetros y geometría, simulación estadística "
"multi-semilla y declaración explícita de supuestos y limitaciones en cada etapa.",
]

RECOMENDACIONES = [
"Las recomendaciones siguientes proyectan el trabajo hacia su implementación y hacia investigaciones futuras; se distinguen de las "
"conclusiones en que no afirman resultados, sino cursos de acción.",

"1. Realizar la campaña de medición radioeléctrica punto a punto en el nivel NV1640 —niveles de señal, dispersión temporal y ancho de banda "
"de coherencia— para calibrar el modelo de propagación con datos propios del área de teleoperación y consolidar el plan de celdas definitivo.",

"2. Obtener cotizaciones vigentes de los fabricantes y ejecutar, junto con el área de planeamiento de la mina, el análisis costo-beneficio "
"cuantitativo (CAPEX/OPEX, costo por metro de galería cubierta, costo por LHD teleoperado y horizonte de recuperación) definido en el "
"objetivo específico 6.",

"3. Implementar un piloto controlado en el corredor de teleoperación con los equipos reales antes de la operación regular, verificando el "
"comportamiento del mesh del fabricante, el traspaso con el vehículo en movimiento y el modo seguro ante pérdida de enlace, video o comandos.",

"4. Definir e implantar, antes de instalar cámaras, la política de tratamiento del video conforme a la Ley N.° 29733: finalidad declarada, "
"acceso restringido, retención mínima y comunicación transparente al personal.",

"5. Desarrollar el plan de capacitación y reconversión del personal operador en paralelo al despliegue técnico, conforme a los compromisos "
"de gestión responsable del cambio asumidos en la sección 1.5.",

"6. Como líneas de investigación futura: extender el modelo a múltiples vehículos y frentes simultáneos (capacidad, interferencia co-canal y "
"reuso de canales); incorporar escenarios de falla de infraestructura e interferencia externa; medir en campo la variabilidad del canal "
"(shadowing) para contrastarla con las bandas del modelo; y evaluar la evolución hacia redes celulares privadas como complemento del acceso "
"802.11ac cuando el caso de negocio lo justifique.",
]

# --- escribir CONCLUSIONES: reemplazar placeholder + insertar el resto antes del heading RECOMENDACIONES
p_reco_heading = d.paragraphs[i_reco]
replace(d.paragraphs[i_conc_ph], CONCLUSIONES[0])
for txt in CONCLUSIONES[1:]:
    p_reco_heading.insert_paragraph_before(txt)

# --- escribir RECOMENDACIONES: reemplazar placeholder + insertar el resto antes de REFERENCIAS
p_refs_heading = d.paragraphs[i_refs + len(CONCLUSIONES) - 1]  # índice corrido por las inserciones
assert p_refs_heading.text.strip() == "REFERENCIAS BIBLIOGRÁFICAS", p_refs_heading.text[:40]
replace(d.paragraphs[i_reco_ph + len(CONCLUSIONES) - 1], RECOMENDACIONES[0])
for txt in RECOMENDACIONES[1:]:
    p_refs_heading.insert_paragraph_before(txt)

d.save(DOC)
print(f"CONCLUSIONES: {len(CONCLUSIONES)} párrafos | RECOMENDACIONES: {len(RECOMENDACIONES)} párrafos")
