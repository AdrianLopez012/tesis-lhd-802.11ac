# -*- coding: utf-8 -*-
import docx, sys
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)
for i in (118, 339, 714, 820):
    p = d.paragraphs[i]
    pb = p.paragraph_format.page_break_before
    prev_xml = d.paragraphs[i-1]._element.xml
    manual = 'type="page"' in prev_xml
    print(f"[{i}] {p.text.strip()[:45]!r} pbb={pb} | salto manual previo={manual}")
