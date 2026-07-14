# -*- coding: utf-8 -*-
# Inserta la figura del patrón de antena al final del Anexo G con caption SEQ
# (numeración automática — tomará el número siguiente al actualizar campos).
import docx, sys
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
IMG = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\patron_antena_3d.png"
d = docx.Document(DOC)

iH = None
for i, p in enumerate(d.paragraphs):
    if p.text.strip().startswith("Anexo H."):
        iH = i; break
assert iH
ref = d.paragraphs[iH]
CAPTION_STYLE = next(p.style for p in d.paragraphs if p.style.name == "Caption")

# texto + imagen + caption SEQ + fuente
t = ref.insert_paragraph_before(
"Como evidencia complementaria de la selección de antena, la figura siguiente presenta el patrón de radiación calculado de una antena "
"helicoidal en modo axial —representativa de la antena embarcada del LHD—: el sólido tridimensional de directividad y los cortes polares "
"de elevación y azimut.")
fp = ref.insert_paragraph_before("")
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run().add_picture(IMG, width=Cm(16.5))
cap = ref.insert_paragraph_before("Figura ")
cap.style = CAPTION_STYLE
fld = cap._element.makeelement(qn("w:fldSimple"), {qn("w:instr"): r" SEQ Figura \* ARABIC "})
rnum = cap._element.makeelement(qn("w:r"), {}); tnum = cap._element.makeelement(qn("w:t"), {})
tnum.text = "19"; rnum.append(tnum); fld.append(rnum)
cap.runs[0]._element.addnext(fld)
cap.add_run(". Patrón de radiación de la antena helicoidal en modo axial: sólido 3D de directividad y cortes de elevación y azimut.")
fu = ref.insert_paragraph_before("Fuente: Elaboración propia con MATLAB Antenna Toolbox (hélice de 7.5 vueltas, 5.0 GHz).")
fu.style = CAPTION_STYLE

# añadir al inventario del Anexo G (última línea "— arquitectura_red...")
for p in d.paragraphs:
    if p.text.strip().startswith("— arquitectura_red.png"):
        p.insert_paragraph_before("— patron_antena_3d.png: Patrón 3D y cortes polares de la antena helicoidal (MATLAB Antenna Toolbox).")
        break
d.save(DOC)
print("patrón insertado en Anexo G con caption SEQ")
