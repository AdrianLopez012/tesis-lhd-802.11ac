# -*- coding: utf-8 -*-
# ETAPA D2 — Tablas de los anexos B/E/H: contenido vigente v3-REAL
import docx, sys
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)

def set_cell(t, r, c, txt):
    cell = t.cell(r, c)
    if cell.paragraphs and cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = txt
        for rr in cell.paragraphs[0].runs[1:]: rr.text = ""
        for pp in cell.paragraphs[1:]:
            for rr in pp.runs: rr.text = ""
    else:
        cell.text = txt
def fill(t, r, vals):
    for c, v in enumerate(vals): set_cell(t, r, c, v)
def del_row(t, r):
    row = t.rows[r]._element; row.getparent().remove(row)

# localizar tablas por contenido (los índices cambiaron)
tB = tE = tX = tH = None
for t in d.tables:
    c0 = t.cell(0,0).text.strip(); c1 = t.cell(1,0).text.strip() if len(t.rows)>1 else ""
    if c0 == "Archivo" and "lhd-teleop" in c1: tB = t
    elif c0 == "Parámetro" and "Tecnología" in c1: tE = t
    elif c0.startswith("Archivo XML"): tX = t
    elif c0 == "Ruta interna": tH = t
assert all(x is not None for x in (tB, tE, tX, tH)), (tB, tE, tX, tH)

# ---- Anexo B: matriz de trazabilidad (19 filas) ----
FILAS_B = [
 ["Archivo","Estado / función"],
 ["cap3/simulacion_ns3/lhd-teleop-v3-real.cc","Simulación principal vigente (Anexo C.1)."],
 ["cap3/simulacion_ns3/geometria_nv1640.h","Geometría real autogenerada (Anexo C.2)."],
 ["cap3/simulacion_ns3/recorrido_nv1640.h","Recorrido real del ciclo autogenerado (Anexo C.3)."],
 ["cap3/simulacion_ns3/parametros_rf.h","Parámetros RF autogenerados (Anexo C.4)."],
 ["cap3/simulacion_ns3/parametros_rf.py","Fuente única de parámetros RF y de propagación (Anexo D.1)."],
 ["graficas_simulacion/mapa_nv1640_datos.py","Fuente única de geometría y posiciones de AP (Anexo D.2)."],
 ["graficas_simulacion/generar_geometria_h.py","Generador geometría → C++ (Anexo D.3)."],
 ["graficas_simulacion/generar_recorrido_h.py","Generador recorrido → C++ (Anexo D.4)."],
 ["graficas_simulacion/generar_parametros_h.py","Generador parámetros → C++ (Anexo D.5)."],
 ["cap3/simulacion_ns3/run_escenarios_v3.sh","Batería de 18 corridas en serie (Anexo E.1)."],
 ["results/principal_s1..s10_v3_*","Escenario de operación, 10 semillas (Anexo F)."],
 ["results/baseline_v3_*","Referencia estática (Anexo F)."],
 ["results/estres_video_v3_* / estres_lhd_v3_*","Escenarios de estrés (Anexo F)."],
 ["results/handover*_v3_*","Escenario dedicado de traspaso, 5 semillas (Anexo F)."],
 ["graficas_simulacion/link_budget.py / .png","Presupuesto de enlace y radios máximos (sección 3.3, Anexo G)."],
 ["graficas_simulacion/mapa_cobertura_pro.py / .png","Cobertura con el mismo modelo de la simulación (Anexo G)."],
 ["graficas_simulacion/resultados_v3.py / kpis_v3real.py","Figuras de resultados del Capítulo 3 (Anexo G)."],
 ["graficas_simulacion/comparacion_kpis.py / .png","Figura 11 del Capítulo 4 (Anexo G)."],
]
while len(tB.rows) > len(FILAS_B): del_row(tB, len(tB.rows)-1)
for i, fila in enumerate(FILAS_B):
    fill(tB, i, fila + [""]*(len(tB.columns)-len(fila)))

# ---- Anexo E: parámetros ejecutados (12 filas) ----
FILAS_E = [
 ["Parámetro","Valor ejecutado","Fuente"],
 ["Tecnología de acceso","IEEE 802.11ac, 5 GHz, 40 MHz, MIMO 2×2","parametros_rf.py"],
 ["Potencia/ganancia Hawk","30 dBm / 11 dBi","Datasheet (parametros_rf.py)"],
 ["Potencia/ganancia Cardinal","23 dBm / 7.5 dBi","Datasheet (parametros_rf.py)"],
 ["Antena del LHD","HELI-40, 4.8 dBi","Datasheet (parametros_rf.py)"],
 ["Modelo de propagación","two-slope n1=1.9, n2=3.4, d_bp=40 m","PROPAGACION (fuente única)"],
 ["Penalización NLOS","10 dB por galería cruzada","PROPAGACION (fuente única)"],
 ["Pérdidas de sistema","9.4 dB","parametros_rf.py"],
 ["Tráfico","Video VBR 40 Mbps (30 fps) + comandos 0.5 + telemetría 0.1","lhd-teleop-v3-real.cc"],
 ["Roaming","Beacon 102.4 ms; MaxMissedBeacons 10 (operación) / 3 (traspaso)","lhd-teleop-v3-real.cc"],
 ["Velocidad LHD","2.22 m/s (operación); 4.0 m/s (estrés)","recorrido_nv1640.h"],
 ["Duración por corrida","300 s; batería de 18 corridas en serie","run_escenarios_v3.sh"],
]
while len(tE.rows) > len(FILAS_E): del_row(tE, len(tE.rows)-1)
for i, fila in enumerate(FILAS_E): fill(tE, i, fila)

# ---- Anexo H: XML FlowMonitor (listar los vigentes) ----
import glob, os
RES = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\cap3\simulacion_ns3\results"
xmls = sorted(os.path.basename(p) for p in glob.glob(os.path.join(RES, "*_v3_flowmon.xml")))
FILAS_X = [["Archivo XML FlowMonitor","Corrida"]] + [[x, x.replace("_v3_flowmon.xml","")] for x in xmls[:len(tX.rows)-1]]
while len(tX.rows) > len(FILAS_X): del_row(tX, len(tX.rows)-1)
for i, fila in enumerate(FILAS_X):
    fill(tX, i, fila + [""]*(len(tX.columns)-len(fila)))

# ---- Anexo H: estructura del anexo digital ----
FILAS_H = [
 ["Ruta interna","Contenido"],
 ["codigo/","lhd-teleop-v3-real.cc y encabezados autogenerados (.h)"],
 ["fuente_unica/","parametros_rf.py y mapa_nv1640_datos.py + generadores"],
 ["scripts_figuras/","Scripts Python de todas las figuras del Anexo G"],
 ["results/","CSV y XML íntegros de las 18 corridas (flow_stats, pos_log, assoc_log, flowmon)"],
 ["figuras/","PNG en resolución completa y animación del recorrido"],
 ["ejecucion/","run_escenarios_v3.sh e instrucciones de reproducción (README)"],
]
while len(tH.rows) > len(FILAS_H): del_row(tH, len(tH.rows)-1)
for i, fila in enumerate(FILAS_H): fill(tH, i, fila)

d.save(DOC)
print(f"tablas B({len(FILAS_B)}f) E({len(FILAS_E)}f) XML({len(FILAS_X)}f) digital({len(FILAS_H)}f) actualizadas")
