# -*- coding: utf-8 -*-
# Correcciones del profesor — parte 1 (mecánicas):
#  C5: estilos con color -> negro (captions, citadestacada, títulos 4/5, hipervínculos)
#  C6: recuadro azul (Citadestacada) -> texto normal
#  C7: captions de Tabla N que están DEBAJO -> moverlas ARRIBA de su tabla
#  C5b: tblHeader (repetir fila de títulos) en todas las tablas
#  C9: RECOMENDACIONES page_break_before
#  C2/C4: headings no-capítulo empiezan arriba (space_before=0)
import docx, sys, re
from docx.oxml.ns import qn
from docx.shared import RGBColor
sys.stdout.reconfigure(encoding="utf-8")

F = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(F)
body = d.element.body

# ---------- C5: estilos a negro ----------
NEGRO_STYLES = ['Descripcin','DescripcinCar','Citadestacada','CitadestacadaCar',
                'TtuloTDC','Ttulo4','Ttulo5','Ttulo4Car','Ttulo5Car',
                'Hipervnculo','Hipervnculovisitado']
n_st = 0
for s in d.styles.element.iter(qn('w:style')):
    sid = s.get(qn('w:styleId'))
    if sid in NEGRO_STYLES:
        rpr = s.find(qn('w:rPr'))
        if rpr is not None:
            c = rpr.find(qn('w:color'))
            if c is not None:
                c.set(qn('w:val'), '000000')
                if c.get(qn('w:themeColor')): del c.attrib[qn('w:themeColor')]
                if c.get(qn('w:themeShade')): del c.attrib[qn('w:themeShade')]
                n_st += 1
print(f"[C5] estilos pasados a negro: {n_st}")

# ---------- C6: recuadro azul -> texto normal ----------
n_cd = 0
for p in d.paragraphs:
    if 'itadestacada' in (p.style.name or '') or (p.style.style_id or '') == 'Citadestacada':
        p.style = d.styles['Normal']
        ppr = p._p.find(qn('w:pPr'))
        if ppr is not None:
            for tag in ('w:pBdr','w:ind'):
                e = ppr.find(qn(tag))
                if e is not None: ppr.remove(e)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0,0,0)
            r.font.italic = False
        n_cd += 1
print(f"[C6] recuadros convertidos a texto normal: {n_cd}")

# ---------- C7: mover captions de tabla que están debajo hacia arriba ----------
def ptext(el):
    if el is None or el.tag != qn('w:p'): return ''
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))

els = list(body)
movidas = 0
for j, el in enumerate(els):
    if el.tag != qn('w:tbl'): continue
    prev = els[j-1] if j > 0 else None
    nxt  = els[j+1] if j+1 < len(els) else None
    if re.match(r'\s*Tabla\b', ptext(prev)): continue           # ya está arriba
    if re.match(r'\s*Tabla\b', ptext(nxt)):
        body.remove(nxt)
        el.addprevious(nxt)                                     # mover encima de la tabla
        movidas += 1
els = list(body)  # refrescar
print(f"[C7] captions de tabla movidas arriba: {movidas}")

# ---------- C5b: tblHeader en la primera fila de todas las tablas ----------
n_th = 0
for tbl in d.tables:
    tr0 = tbl.rows[0]._tr
    trPr = tr0.find(qn('w:trPr'))
    if trPr is None:
        trPr = tr0.makeelement(qn('w:trPr'), {})
        tr0.insert(0, trPr)
    if trPr.find(qn('w:tblHeader')) is None:
        th = trPr.makeelement(qn('w:tblHeader'), {})
        trPr.append(th)
        n_th += 1
print(f"[C5b] tablas con fila de titulos repetible añadida: {n_th}")

# ---------- C9: RECOMENDACIONES en página aparte ----------
for p in d.paragraphs:
    if p.text.strip() == 'RECOMENDACIONES Y OBSERVACIONES':
        p.paragraph_format.page_break_before = True
        print("[C9] RECOMENDACIONES: page_break_before = True")
        break

# ---------- C2/C4: headings no-capítulo empiezan ARRIBA ----------
ARRIBA = ['RESUMEN','ABSTRACT','ÍNDICE GENERAL','ÍNDICE DE TABLAS',
          'ÍNDICE DE FIGURAS','ÍNDICE DE ECUACIONES','GLOSARIO','INTRODUCCIÓN']
n_up = 0
ps = d.paragraphs
for i, p in enumerate(ps):
    if p.text.strip() in ARRIBA and p.style.name == 'Heading 1':
        p.paragraph_format.space_before = 0
        n_up += 1
        # quitar párrafos vacíos inmediatamente anteriores (empujan hacia abajo)
        j = i-1
        while j >= 0 and not ps[j].text.strip() and ps[j].style.name == 'Normal':
            pj = ps[j]._p
            ppr = pj.find(qn('w:pPr'))
            tiene_sect = ppr is not None and ppr.find(qn('w:sectPr')) is not None
            if tiene_sect: break                                 # nunca borrar un salto de sección
            pj.getparent().remove(pj)
            j -= 1
print(f"[C2/C4] headings ajustados para empezar arriba: {n_up}")

d.save(F)
print("GUARDADO OK")
