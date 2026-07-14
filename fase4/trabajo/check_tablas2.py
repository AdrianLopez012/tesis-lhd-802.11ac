# -*- coding: utf-8 -*-
# Diagnóstico dirigido: Tablas 21/22/23 del Cap. 3 — columnas vacías,
# filas sobrantes y anchos reales
import docx, sys
sys.stdout.reconfigure(encoding="utf-8")
DOC = r"C:\Users\Adrian Lopez\Documents\tesis_proyecto\fase4\trabajo\tesis_v2.docx"
d = docx.Document(DOC)

for ti in (22, 23, 24):
    t = d.tables[ti]
    ncols = len(t.columns); nrows = len(t.rows)
    print(f"===== tabla idx {ti}: {nrows} filas x {ncols} columnas =====")
    # columnas vacías (todas las celdas sin texto)
    for c in range(ncols):
        vals = [t.cell(r, c).text.strip() for r in range(nrows)]
        if all(not v for v in vals):
            print(f"  >>> COLUMNA {c} completamente VACÍA")
    # contenido resumido
    for r in range(nrows):
        celdas = [t.cell(r, c).text.strip()[:30] for c in range(ncols)]
        print(f"  f{r}: {celdas}")
    w = sum((c.width or 0) for c in t.rows[0].cells)
    print(f"  ancho total: {w/914400:.2f} in\n")
